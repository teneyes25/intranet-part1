# -*- coding: utf-8 -*-
"""Ten Eyes Intranet Part 1 — PRD v1.2.1 (.docx) generator."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "Ten_Eyes_Intranet_Part1_PRD_v1.2.1.docx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GRAY = RGBColor(0x47, 0x55, 0x69)
LIGHT_GRAY = RGBColor(0x64, 0x74, 0x8B)


def set_run_font(run, size=11, bold=False, color=None, name="맑은 고딕"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, color=NAVY)
    return p


def add_para(doc, text, size=11, bold=False, color=GRAY, space_after=8, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        p.alignment = align
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11, color=GRAY)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_screenshot_placeholder(doc, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[ 스크린샷 자리: {caption} ]")
    set_run_font(run, size=10, color=LIGHT_GRAY, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    # light border-like note
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = note.add_run("※ 배포 후 실제 화면 캡처를 이 위치에 삽입하세요.")
    set_run_font(r2, size=9, color=LIGHT_GRAY)
    note.paragraph_format.space_after = Pt(12)


def set_cell_text(cell, text, bold=False, size=10, color=GRAY):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def shade_header_row(row):
    from docx.oxml import OxmlElement

    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "0F172A")
        tcPr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.bold = True


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
    shade_header_row(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val, size=10)
    doc.add_paragraph()
    return table


def build():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # ========== COVER ==========
    for _ in range(3):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Ten Eyes Intranet Part 1")
    set_run_font(r, size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Product Requirements Document (PRD)")
    set_run_font(r, size=18, bold=True, color=BLUE)

    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ver.add_run("v1.2.1  ·  2026년 7월")
    set_run_font(r, size=14, color=GRAY)

    doc.add_paragraph()

    meta_lines = [
        "프로젝트명: Ten Eyes Intranet (인트라넷 Part 1)",
        "목적: 본사 + 전국 100개 대리점 업무 효율화 및 협업 강화",
        "배포 형태: Docker + Nginx (온프레미스)",
        "상태: Phase 5 완료 — Firebase 설정 제외 상용화 MVP",
        "문서 분류: 제품 요구사항 정의서 (내부용)",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, size=11, color=GRAY)

    doc.add_page_break()

    # ========== TOC (manual) ==========
    add_heading_styled(doc, "목차", 1)
    toc = [
        "1. 프로젝트 개요",
        "2. 프로젝트 배경 및 목표",
        "3. 주요 기능",
        "    3.1 핵심 인트라넷 기능",
        "    3.2 Page Builder (상용화 핵심)",
        "    3.3 기술 스택",
        "4. 기능 상세 설명",
        "5. 사용자 역할 및 권한",
        "6. 비기능 요구사항",
        "7. 기술 스택 다이어그램",
        "8. 배포 및 운영",
        "9. 향후 개발 로드맵 (Part 2)",
        "10. 배포 및 유지보수 가이드 요약",
        "11. 부록 — 문서 맵 및 용어",
    ]
    for item in toc:
        add_para(doc, item, size=11, space_after=4)
    doc.add_page_break()

    # ========== 1. 프로젝트 개요 ==========
    add_heading_styled(doc, "1. 프로젝트 개요", 1)

    add_heading_styled(doc, "1.1 기본 정보", 2)
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["프로젝트명", "Ten Eyes Intranet (인트라넷 Part 1)"],
            ["버전", "v1.2.1 (2026년 7월)"],
            ["개발 기간", "Day 1 ~ Day 36+ (Phase 1~5)"],
            ["배포 형태", "Docker + Nginx 온프레미스 (포트 8080)"],
            ["대상 조직", "본사 + 전국 약 100개 대리점"],
            ["현재 상태", "Phase 5 완료 — Firebase 설정 제외 상용화 MVP"],
            ["기본 데이터 계층", "PostgreSQL + Prisma (Firebase는 선택 활성화)"],
        ],
    )

    add_heading_styled(doc, "1.2 제품 한 줄 정의", 2)
    add_para(
        doc,
        "Ten Eyes Intranet Part 1은 본사와 전국 대리점이 동일한 웹 환경에서 "
        "공지·자료·AI 지원·페이지 커스터마이징을 수행하고, Part 2 B2B 거래 기능을 "
        "안전하게 확장할 수 있도록 설계된 온프레미스 인트라넷 플랫폼이다.",
    )

    add_heading_styled(doc, "1.3 범위 (In Scope / Out of Scope)", 2)
    add_para(doc, "In Scope (Part 1)", bold=True, color=NAVY)
    for t in [
        "인증·RBAC·감사 로그 기반 인트라넷 코어",
        "게시판(공지, 일정, 업계 뉴스, 홍보실, 상품자료실) 및 매장 찾기",
        "AI Agent (RAG), 알림, 대리점 포털, Design System v4 UI",
        "GrapesJS Page Builder (본사/대리점 Self-Service)",
        "회사 테마(French Maison), Docker 배포·운영 문서",
        "Firebase 하이브리드 코드 준비 (기본 FIREBASE_ENABLED=false)",
    ]:
        add_bullet(doc, t)

    add_para(doc, "Out of Scope (Part 2 이후)", bold=True, color=NAVY)
    for t in [
        "실거래 발주·세금계산서·카드결제·전자계약 (현재는 플레이스홀더 블록만)",
        "대리점 간 물품 거래·재고 실운영 모듈",
        "PG 연동, 전자계약 서비스 연동",
        "Firebase Console 실서비스 설정 (배포 후 선택)",
    ]:
        add_bullet(doc, t)

    # ========== 2. 배경 및 목표 ==========
    add_heading_styled(doc, "2. 프로젝트 배경 및 목표", 1)

    add_heading_styled(doc, "2.1 배경", 2)
    add_para(
        doc,
        "전국 대리점 네트워크를 운영하는 조직에서는 공지·상품 자료·홍보 콘텐츠가 "
        "채널별로 분산되기 쉽고, 대리점별 브랜딩·레이아웃 요구가 본사 IT에 집중되는 "
        "병목이 발생한다. 또한 업계 뉴스·사내 자료를 빠르게 검색·요약하는 AI 지원과, "
        "향후 B2B 거래(발주·결제·계약)로 확장 가능한 단일 플랫폼이 필요하다.",
    )
    add_para(
        doc,
        "Part 1은 이러한 문제를 해결하기 위해 PostgreSQL 중심의 안정적인 온프레미스 "
        "인트라넷을 먼저 상용화하고, Page Builder로 대리점 Self-Service를 제공하며, "
        "Firebase·거래 모듈은 하이브리드/플레이스홀더로 준비해 단계적 확장을 가능하게 한다.",
    )

    add_heading_styled(doc, "2.2 목표", 2)
    goals = [
        ("업무 효율화", "공지·자료·일정·뉴스를 단일 포털에서 조회·검색·공유"),
        ("협업 강화", "본사–대리점 역할 기반 접근과 감사 로그로 투명한 운영"),
        ("Self-Service", "대리점이 자사 레이아웃·브랜딩을 직접 편집 (Page Builder)"),
        ("AI 지원", "인트라넷 자료 RAG + 업계 뉴스 자동 수집으로 정보 접근성 향상"),
        ("온프레미스 상용화", "Docker·Nginx·체크리스트 기반 본사 서버 배포"),
        ("확장 준비", "Part 2 거래 API(/api/v2) 및 Firebase 선택 활성화 구조"),
    ]
    add_table(doc, ["목표", "설명"], [[a, b] for a, b in goals])

    add_heading_styled(doc, "2.3 성공 기준 (MVP)", 2)
    for t in [
        "Credentials 로그인 및 ADMIN / HEADQUARTER / DEALER RBAC 정상 동작",
        "상용 게시판 메뉴(공지·일정·업계 뉴스·홍보실·상품자료실·매장 찾기) 제공",
        "Page Builder: ADMIN 저장, HQ read-only, DEALER 자사 scope 잠금",
        "FIREBASE_ENABLED=false 상태에서 PostgreSQL·로컬 업로드만으로 운영 가능",
        "production-check / smoke / typecheck / lint 통과",
        "본사 배포 체크리스트(HQ_DEPLOYMENT_FINAL_CHECKLIST) 문서화 완료",
    ]:
        add_bullet(doc, t)

    # ========== 3. 주요 기능 ==========
    add_heading_styled(doc, "3. 주요 기능", 1)

    add_heading_styled(doc, "3.1 핵심 인트라넷 기능", 2)

    add_para(doc, "인증 및 접근 제어", bold=True, color=NAVY)
    for t in [
        "회원가입 / 로그인: NextAuth Credentials (이메일·비밀번호)",
        "Google 로그인: Firebase Auth 연동 (FIREBASE_ENABLED=true 시)",
        "Role-based Access Control: ADMIN, HEADQUARTER, DEALER",
        "미들웨어 기반 라우트 보호 및 API 권한 검증",
    ]:
        add_bullet(doc, t)

    add_para(doc, "게시판 및 콘텐츠", bold=True, color=NAVY)
    for t in [
        "공지사항, 일정 게시판, 업계 뉴스, 홍보실(방송 자료실), 상품 자료실",
        "CRUD, 검색, 파일 첨부 (로컬 uploads 볼륨 / Firebase Storage 폴백)",
        "매장 찾기 (/stores)",
        "업계 뉴스 자동 수집: 관리자 키워드 설정, 매일 09:00 KST cron",
        "메뉴 숨김 토글(nav-modules): 문서함·근태·전자결재·통합검색 — 라우트 유지, 메뉴만 off",
    ]:
        add_bullet(doc, t)

    add_para(doc, "AI Agent", bold=True, color=NAVY)
    for t in [
        "RAG 기반 Q&A: 인트라넷 자료 벡터 검색(pgvector) + Full-Text Search",
        "RRF 리랭크 후 OpenAI(gpt-4o-mini) 답변, 추출형 폴백",
        "대화 메모리(ai_conversations), 출처(Markdown Sources) 표시",
        "업계 뉴스 수집 데이터와 연계 가능한 학습·검색 파이프라인",
    ]:
        add_bullet(doc, t)

    add_para(doc, "업무 모듈 (코드 유지, 기본 메뉴 숨김)", bold=True, color=NAVY)
    for t in [
        "전자결재: 기안, 결재선, AI 지원 (moduleKey: approval)",
        "근태관리, 문서함, 통합 검색 — 필요 시 NAV_MODULE_VISIBLE로 재노출",
        "알림 시스템: DB 알림 + 탑바 벨 (Firebase Firestore/FCM은 선택)",
    ]:
        add_bullet(doc, t)

    add_para(doc, "대리점·운영", bold=True, color=NAVY)
    for t in [
        "대리점 포털: 주문·재고·공지 KPI (실데이터 연동)",
        "거래 모듈 허브(/commerce): Part 2 Preview 페이지",
        "관리: 사용자, 활동 로그, 모니터링, 커스터마이징, 뉴스 키워드",
        "회사 테마(French Maison): pageBuilderConfig JSON, DB 스키마 변경 없음",
        "Design System v4: Navy/Blue 토큰, PageHero, KPI, shadcn/ui",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "3.2 Page Builder (상용화 핵심)", 2)
    add_para(
        doc,
        "GrapesJS 기반 Visual Page Builder는 Part 1 상용화의 핵심 차별화 기능이다. "
        "관리자는 본사(global) 및 전 대리점 레이아웃을 편집하고, 대리점은 자사 scope만 "
        "Self-Service로 편집한다. HEADQUARTER는 미리보기 전용(read-only)이다.",
    )
    for t in [
        "관리자용 Visual Editor: /admin/customize/page-builder",
        "대리점 Self-Service: dealerId 자동 잠금, 자사 레이아웃만 저장",
        "템플릿 갤러리: 9+ built-in, HTML 썸네일, projectData+HTML 동시 저장",
        "Undo/Redo (최대 64단계), 버전 관리 (저장 시 백업, 최대 3버전)",
        "Slug: dashboard, banner, board:{key}, post-detail:{boardKey|postId}",
        "게시판/상품 상세 페이지 레이아웃 커스터마이징",
        "기기 미리보기: 데스크톱 / 태블릿 / 모바일",
        "HTML sanitize, 감사 로그(page_layout:{slug}:{dealerId})",
        "거래 블록(발주·세금계산서·결제·계약)은 Part 2 플레이스홀더만 제공",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "3.3 기술 스택", 2)
    add_table(
        doc,
        ["계층", "기술"],
        [
            ["Frontend", "Next.js 15 (App Router), TypeScript, TailwindCSS, shadcn/ui"],
            ["Backend", "Next.js API Routes, Middleware RBAC"],
            ["Database", "PostgreSQL + Prisma, pgvector (RAG)"],
            ["Cache", "Redis"],
            ["AI", "OpenAI + LangChain(text splitters), Hybrid Search"],
            ["Page Builder", "GrapesJS"],
            ["Auth", "NextAuth.js + (선택) Firebase Auth"],
            ["Real-time / Push", "(선택) Firebase Firestore, FCM"],
            ["Storage", "로컬 uploads 볼륨 + (선택) Firebase Storage"],
            ["Container", "Docker Compose + Nginx (:8080)"],
            ["테스트", "Playwright E2E, production-check, smoke, load-test"],
        ],
    )

    # ========== 4. 기능 상세 ==========
    add_heading_styled(doc, "4. 기능 상세 설명", 1)
    add_para(
        doc,
        "아래 각 기능 블록에는 배포 후 실제 UI 스크린샷을 삽입할 자리를 표시한다. "
        "문서 검토·고객 설명 시 캡처를 교체하면 된다.",
    )

    features_detail = [
        (
            "4.1 로그인 및 인증",
            [
                "이메일/비밀번호 Credentials 로그인 (기본 운영 경로)",
                "세션: NextAuth JWT, PostgreSQL User와 연동",
                "FIREBASE_ENABLED=true 시 Google 로그인 버튼 노출",
                "데모 계정: admin@intranet.local, hq@intranet.local, dealer@intranet.local",
            ],
            "로그인 화면 (AuthShell)",
        ),
        (
            "4.2 대시보드",
            [
                "PageHero + KPI strip, Design System v4 적용",
                "Page Builder dashboard slug 레이아웃 오버레이",
                "역할별 바로가기 및 위젯",
            ],
            "대시보드 홈",
        ),
        (
            "4.3 게시판",
            [
                "공지사항 / 일정 / 업계 뉴스 / 홍보실 / 상품자료실",
                "목록·상세·작성·수정·삭제, 검색, 첨부파일",
                "board:{key} · post-detail 레이아웃 커스터마이징",
            ],
            "게시판 목록 및 상세",
        ),
        (
            "4.4 AI 에이전트",
            [
                "자연어 질문 → Intent → Hybrid Search → LLM 답변",
                "출처 문서 링크, 대화 이력 유지",
                "OpenAI API 키 필요 (미설정 시 폴백/제한 동작)",
            ],
            "AI Agent 채팅 화면",
        ),
        (
            "4.5 Page Builder",
            [
                "블록 패널: 히어로, 배너, 카드, AI 위젯, 공지, 갤러리 등",
                "템플릿 적용 → 편집 → 저장 → 버전 복원",
                "ADMIN / DEALER / HEADQUARTER 권한 분기",
            ],
            "GrapesJS Visual Editor",
        ),
        (
            "4.6 대리점 포털",
            [
                "DEALER·ADMIN 접근",
                "주문·재고·공지 KPI, 빠른 작업 그리드",
                "페이지 편집 진입점 (도구 → 페이지 편집)",
            ],
            "대리점 포털",
        ),
        (
            "4.7 관리자 기능",
            [
                "커스터마이징(브랜딩·Page Builder), 뉴스 키워드",
                "활동 로그(/admin/activity), 모니터링(/admin/monitoring)",
                "사용자·대리점 관리",
            ],
            "관리자 커스터마이징 센터",
        ),
    ]

    for title, bullets, shot in features_detail:
        add_heading_styled(doc, title, 2)
        for b in bullets:
            add_bullet(doc, b)
        add_screenshot_placeholder(doc, shot)

    # ========== 5. 역할 ==========
    add_heading_styled(doc, "5. 사용자 역할 및 권한", 1)
    add_para(
        doc,
        "역할은 Prisma Role enum(ADMIN, HEADQUARTER, DEALER)으로 정의되며, "
        "네비게이션·API·Page Builder 저장 권한에 일관 적용된다.",
    )
    add_table(
        doc,
        ["역할", "설명", "주요 권한", "Page Builder"],
        [
            [
                "ADMIN",
                "시스템 관리자",
                "전체 관리, 브랜딩, 모니터링, 사용자·대리점",
                "global + 전 대리점 저장",
            ],
            [
                "HEADQUARTER",
                "본사 사용자",
                "본사 게시판·업무 조회, 홍보실 등",
                "미리보기 전용 (저장 차단)",
            ],
            [
                "DEALER",
                "대리점 사용자",
                "자사 포털, 게시판 조회, 자사 데이터",
                "자사 dealerId만 저장",
            ],
        ],
    )
    add_para(doc, "권한 매트릭스 (요약)", bold=True, color=NAVY)
    add_table(
        doc,
        ["기능", "ADMIN", "HEADQUARTER", "DEALER"],
        [
            ["로그인 / 대시보드", "O", "O", "O"],
            ["공지·일정·뉴스·상품자료실", "O", "O", "O"],
            ["홍보실", "O", "O", "제한"],
            ["대리점 포털", "O", "—", "O"],
            ["AI 에이전트", "O", "O", "O"],
            ["관리자 메뉴", "O", "—", "—"],
            ["Page Builder 저장", "O", "X (read-only)", "자사만"],
            ["감사 로그 / 모니터링", "O", "—", "—"],
        ],
    )

    # ========== 6. 비기능 ==========
    add_heading_styled(doc, "6. 비기능 요구사항", 1)

    add_heading_styled(doc, "6.1 보안", 2)
    for t in [
        "RBAC: 미들웨어 + API assert (역할·dealerId scope)",
        "Rate Limiting, CSP, 보안 헤더 (X-Frame-Options 등)",
        "Audit Log: 주요 관리·Page Builder 변경 기록",
        "HTML sanitize (Page Builder 렌더)",
        "시크릿: NEXTAUTH_SECRET, POSTGRES_PASSWORD — .env.production (git 제외)",
        "DB 포트(5432)·Redis(6379) 외부 차단, HTTP 8080만 사내망 개방 권장",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "6.2 성능", 2)
    for t in [
        "Next.js standalone Docker 이미지, Nginx 리버스 프록시",
        "Page Builder 에디터 1회 초기화 + slug/scope 증분 로드",
        "load-test.mjs, production-check, smoke 스크립트",
        "Redis 캐시 계층",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "6.3 가용성·운영", 2)
    for t in [
        "온프레미스 Docker Compose 단일 스택 배포",
        "헬스체크: GET /api/health",
        "백업: PostgreSQL pg_dump + uploads 볼륨",
        "모니터링: /admin/monitoring",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "6.4 확장성", 2)
    for t in [
        "Firebase 하이브리드: Auth/Storage/Firestore/FCM — 플래그로 on/off",
        "Page Builder slug·scope로 페이지·대리점 단위 확장",
        "Part 2 API는 /api/v2 네임스페이스 (Part 1 /api/* breaking change 없음)",
        "nav-modules로 메뉴 단계적 공개",
    ]:
        add_bullet(doc, t)

    # ========== 7. 다이어그램 ==========
    add_heading_styled(doc, "7. 기술 스택 다이어그램", 1)
    add_para(
        doc,
        "아래는 텍스트 기반 아키텍처 다이어그램이다. 상세 Mermaid 원본은 docs/ARCHITECTURE.md를 참고한다.",
    )

    add_heading_styled(doc, "7.1 시스템 개요", 2)
    diagram = """
[브라우저 / 모바일·태블릿]
            │
            ▼
     [Nginx :8080]
            │
            ▼
┌───────────────────────────────────────┐
│     Next.js 15 App (standalone)       │
│  App Router · API Routes · Middleware │
│  AI Layer (RAG) · Page Builder API    │
└───────────────────────────────────────┘
     │            │            │
     ▼            ▼            ▼
[PostgreSQL]  [Redis]   [uploads volume]
 + pgvector
     │
     └──── (선택) Firebase Auth / Storage / Firestore / FCM
"""
    p = doc.add_paragraph()
    run = p.add_run(diagram.strip("\n"))
    set_run_font(run, size=9, name="Consolas", color=NAVY)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    add_heading_styled(doc, "7.2 AI / RAG 파이프라인", 2)
    ai_diagram = """
User Question → Intent Detection → Hybrid Search
                                      ├─ Vector (pgvector)
                                      └─ Full-Text Search
                                              ↓
                                         RRF Rerank
                                      ├─ OpenAI (gpt-4o-mini)
                                      └─ Extractive Fallback
                                              ↓
                              Markdown Answer + Sources
                                              ↓
                                    ai_conversations (memory)
"""
    p = doc.add_paragraph()
    run = p.add_run(ai_diagram.strip("\n"))
    set_run_font(run, size=9, name="Consolas", color=NAVY)

    add_heading_styled(doc, "7.3 Firebase 하이브리드 (선택)", 2)
    add_para(
        doc,
        "기본값 FIREBASE_ENABLED=false. 활성화 시에만 Google 로그인, Storage, "
        "Firestore 실시간(알림·Page Builder 협업), FCM이 동작하며, 미설정 시 "
        "Credentials·로컬 uploads·알림 폴링으로 100% 동작한다.",
    )

    # ========== 8. 배포 ==========
    add_heading_styled(doc, "8. 배포 및 운영", 1)

    add_heading_styled(doc, "8.1 배포 형태", 2)
    for t in [
        "Docker Compose: app + PostgreSQL + Redis + Nginx",
        "접속: http://<본사서버IP>:8080",
        "명령: docker compose --env-file .env.production up -d --build",
        "또는 개발 PC: npm run docker:sync",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "8.2 필수 환경변수 (.env.production)", 2)
    add_table(
        doc,
        ["변수", "설명"],
        [
            ["POSTGRES_PASSWORD", "DB 비밀번호 (강력 랜덤)"],
            ["NEXTAUTH_SECRET", "세션 시크릿 (openssl rand -base64 32)"],
            ["NEXTAUTH_URL", "http://<본사IP>:8080"],
            ["ALLOWED_ORIGINS", "동일 origin"],
            ["RUN_SEED", "운영 false"],
            ["FIREBASE_ENABLED", "false (기본)"],
            ["OPENAI_API_KEY", "AI 사용 시 (선택)"],
            ["NAVER_CLIENT_ID/SECRET", "업계 뉴스 수집 시 (선택)"],
        ],
    )

    add_heading_styled(doc, "8.3 Firebase 선택적 활성화", 2)
    for t in [
        "Firebase Console 프로젝트·Auth·Firestore·Storage·FCM 설정",
        "NEXT_PUBLIC_FIREBASE_* 및 FIREBASE_SERVICE_ACCOUNT_JSON 설정",
        "FIREBASE_ENABLED=true 후 재빌드",
        "상세: docs/FIREBASE_HYBRID.md",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "8.4 백업·모니터링", 2)
    for t in [
        "DB: pg_dump 또는 volume 스냅샷 정기 백업",
        "파일: uploads 볼륨 백업",
        "모니터링: /admin/monitoring, /api/health",
        "감사: /admin/activity",
        "롤백: 이전 Docker 이미지 + DB 복원",
    ]:
        add_bullet(doc, t)

    # ========== 9. 로드맵 ==========
    add_heading_styled(doc, "9. 향후 개발 로드맵 (Part 2)", 1)
    add_para(
        doc,
        "Part 2는 B2B 거래·결제·계약 도메인을 /api/v2로 추가하며, Part 1 API와 "
        "스키마를 깨지 않는 방향으로 확장한다. Page Builder의 거래 블록은 "
        "실모듈과 하이드레이션으로 연결한다.",
    )
    add_table(
        doc,
        ["단계", "내용", "비고"],
        [
            ["Part 2-A", "발주서(Order) CRUD, 대리점 여신·사업자 정보", "/api/v2/orders"],
            ["Part 2-B", "세금계산서, 카드결제(PG 연동)", "Payment 모듈"],
            ["Part 2-C", "표준계약서·전자계약", "Contract 모듈"],
            ["Part 2-D", "대리점 간 물품 거래, 재고 관리", "Shipment·Inventory"],
            ["Part 2-E", "Firebase 실서비스 활성화 (선택)", "Auth/Push/Storage"],
            ["지속", "UI/UX·성능·보안 고도화", "Design System 유지"],
        ],
    )
    add_para(doc, "Part 2 RBAC (초안)", bold=True, color=NAVY)
    add_table(
        doc,
        ["역할", "주문", "결제", "배송 상태"],
        [
            ["DEALER", "본인 대리점 CRUD", "본인", "조회만"],
            ["HEADQUARTER", "전체 조회", "전체 조회", "상태 변경"],
            ["ADMIN", "전체", "전체", "전체"],
        ],
    )

    # ========== 10. 유지보수 가이드 ==========
    add_heading_styled(doc, "10. 배포 및 유지보수 가이드 요약", 1)

    add_heading_styled(doc, "10.1 배포 전 검증", 2)
    for t in [
        "npm run typecheck",
        "npm run lint",
        "npm run build",
        "npm run production:check -- http://localhost:8080",
        "npm run smoke -- http://localhost:8080",
        "npm run test:ui-upgrade (선택: npm run test:e2e)",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "10.2 배포 직후", 2)
    for t in [
        "GET /api/health → healthy",
        "관리자 로그인 후 비밀번호 즉시 변경",
        "/admin/customize — 앱명·로고·브랜딩",
        "/admin/customize/page-builder — 대시보드·배너 저장",
        "/admin/news-keywords — 업계 뉴스 키워드",
        "최초 1회: npm run db:seed-company-theme, db:seed-boards-nav",
        "대리점 1곳 Pilot — DEALER Page Builder 저장 테스트",
    ]:
        add_bullet(doc, t)

    add_heading_styled(doc, "10.3 일상 운영", 2)
    for t in [
        "백업 스케줄 확인 (DB + uploads)",
        "활동 로그·모니터링 주기 점검",
        "뉴스 수집 cron(09:00 KST) 및 API 키 유효성",
        "메뉴 재노출이 필요하면 nav-modules.ts 토글",
        "장애 시: docker compose logs, health, 롤백 절차",
    ]:
        add_bullet(doc, t)

    # ========== 11. 부록 ==========
    add_heading_styled(doc, "11. 부록 — 문서 맵 및 용어", 1)

    add_heading_styled(doc, "11.1 관련 문서", 2)
    add_table(
        doc,
        ["문서", "용도"],
        [
            ["docs/HQ_DEPLOYMENT_FINAL_CHECKLIST.md", "본사 서버 배포 최종 체크리스트"],
            ["docs/MVP_DEPLOYMENT_CHECKLIST.md", "MVP 배포 요약"],
            ["DEPLOY.md", "상세 배포 가이드"],
            ["docs/PAGE_BUILDER.md", "Page Builder 운영"],
            ["docs/FIREBASE_HYBRID.md", "Firebase 활성화"],
            ["docs/ARCHITECTURE.md", "아키텍처 Mermaid"],
            ["docs/PART2_API_DESIGN.md", "Part 2 API 초안"],
            ["CHANGELOG.md", "버전별 변경 이력"],
            ["ADMIN_MANUAL.md / USER_MANUAL.md", "관리자·사용자 매뉴얼"],
        ],
    )

    add_heading_styled(doc, "11.2 용어", 2)
    add_table(
        doc,
        ["용어", "정의"],
        [
            ["Part 1", "인트라넷 코어·Page Builder·AI (본 문서 범위)"],
            ["Part 2", "B2B 거래·결제·계약·재고 확장"],
            ["global scope", "본사 공통 Page Builder 레이아웃 (dealerId=global)"],
            ["Self-Service", "대리점이 자사 레이아웃을 직접 편집하는 권한 모델"],
            ["RAG", "Retrieval-Augmented Generation — 검색 증강 생성"],
            ["상용화 MVP", "Firebase Console 설정 없이 PostgreSQL만으로 운영 가능한 v1.2.1"],
        ],
    )

    add_heading_styled(doc, "11.3 문서 이력", 2)
    add_table(
        doc,
        ["버전", "일자", "내용"],
        [
            ["1.2.1", "2026-07", "Phase 5 상용화 MVP 기준 최초 PRD 작성"],
        ],
    )

    # Footer note
    doc.add_paragraph()
    add_para(
        doc,
        "— End of Document — Ten Eyes Intranet Part 1 PRD v1.2.1 —",
        size=10,
        color=LIGHT_GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    build()
